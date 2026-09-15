"""입력 파일을 (위치라벨, 텍스트) 단위로 읽고, 보호 구간을 가린 뒤 문장으로 나눈다."""
import os
import re

FENCE = re.compile(r"```.*?```", re.S)
INLINE_CODE = re.compile(r"`[^`\n]+`")
URL = re.compile(r"https?://\S+")
DQUOTE = re.compile(r"\"[^\"\n]{2,}\"")
FRONTMATTER = re.compile(r"\A---\n.*?\n---\n", re.S)
NUM_UNIT = re.compile(r"\d[\d,.]*\s*(?:%|원|달러|명|개|건|시간|분|초|년|월|일|kg|g|km|m|GB|MB|회|배|점|자|줄|쪽|p)\b")
MD_HEADING = re.compile(r"^\s{0,3}#{1,6}\s+(.*)$")
BULLET = re.compile(r"^\s*(?:[-*+]|\d+[.)])\s+")
SENT_END = re.compile(r"(?<=[.!?。])\s+|(?<=[다요죠음임까])\s*\n")
EMOJI = re.compile("[\U0001F300-\U0001FAFF☀-➿⭐✅❌]")


def read_units(path):
    """파일 확장자별로 (라벨, 텍스트) 목록을 돌려준다. pptx는 발표자 노트를 반드시 포함한다."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pptx":
        from pptx import Presentation
        out = []
        for i, slide in enumerate(Presentation(path).slides, 1):
            seen = 0
            for sh in slide.shapes:
                if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
                    seen += 1
                    out.append(("S%d %s" % (i, "제목" if seen == 1 else "본문"), sh.text_frame.text))
            try:
                n = slide.notes_slide.notes_text_frame.text
                if n.strip():
                    out.append(("S%d 노트" % i, n))
            except Exception:
                pass
        return out
    if ext == ".docx":
        import docx
        d = docx.Document(path)
        return [("p%d" % i, p.text) for i, p in enumerate(d.paragraphs, 1) if p.text.strip()]
    with open(path, encoding="utf-8", errors="replace") as f:
        return [("본문", f.read())]


def mask_protected(text, protect):
    """보호 구간을 같은 길이의 공백으로 바꿔 위치는 유지한 채 판정에서 뺀다."""
    def blank(m):
        return re.sub(r"\S", " ", m.group(0))
    if "frontmatter" in protect:
        text = FRONTMATTER.sub(blank, text)
    if "fenced_code" in protect:
        text = FENCE.sub(blank, text)
    if "inline_code" in protect:
        text = INLINE_CODE.sub(blank, text)
    if "urls" in protect:
        text = URL.sub(blank, text)
    if "quotes_double" in protect:
        text = DQUOTE.sub(blank, text)
    if "numbers_units" in protect:
        text = NUM_UNIT.sub(blank, text)
    return text


def lines(text):
    return text.split("\n")


def is_heading(line):
    return bool(MD_HEADING.match(line))


def heading_text(line):
    m = MD_HEADING.match(line)
    return m.group(1).strip() if m else ""


def strip_markup(line):
    line = MD_HEADING.sub(lambda m: m.group(1), line)
    line = BULLET.sub("", line)
    line = re.sub(r"\*\*([^*]+)\*\*", r"\1", line)
    return line.strip()


def paragraphs(text):
    """빈 줄로 나눈 문단. 헤딩, 표, 구분선은 뺀다."""
    out = []
    for block in re.split(r"\n\s*\n", text):
        b = block.strip()
        if not b or is_heading(b) or b.startswith("|") or re.match(r"^[-=*_]{3,}$", b):
            continue
        out.append(b)
    return out


def sentences(text):
    """산문 문장 목록. 표, 헤딩, 코드, 짧은 조각은 뺀다."""
    out = []
    for p in paragraphs(text):
        for raw in p.split("\n"):
            raw = strip_markup(raw)
            if not raw or raw.startswith("|"):
                continue
            for s in SENT_END.split(raw):
                s = s.strip()
                if len(s) >= 6:
                    out.append(s)
    return out
